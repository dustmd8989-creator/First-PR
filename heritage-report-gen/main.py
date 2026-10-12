#!/usr/bin/env python3
"""문화유산 조사보고서 초안 생성기.

조사자 메모(input/) + 보고서 유형(templates/) + 절별 프롬프트(prompts/)
-> Claude가 절별 초안 생성 -> output/*.docx

사용 예:
  python main.py -i input/sample_memo.md -t architecture_survey --dry-run
  ANTHROPIC_API_KEY=... python main.py -i input/sample_memo.md -t architecture_survey
"""
import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).parent
DEFAULT_MODEL = "claude-sonnet-5-5"
BODY_FONT = "바탕"


def load_template(name: str) -> dict:
    path = ROOT / "templates" / f"{name}.json"
    if not path.exists():
        avail = ", ".join(p.stem for p in (ROOT / "templates").glob("*.json"))
        sys.exit(f"템플릿 없음: {name} (사용 가능: {avail})")
    return json.loads(path.read_text(encoding="utf-8"))


def build_prompt(section: dict, memo: str) -> str:
    guide = (ROOT / "prompts" / section["prompt"]).read_text(encoding="utf-8")
    return f"{guide}\n\n# 조사자 메모\n{memo}\n\n위 절의 본문만 작성하라."


def generate(client, model: str, system: str, user: str) -> str:
    resp = client.messages.create(
        model=model,
        max_tokens=2000,
        system=system,
        messages=[{"role": "user", "content": user}],
    )
    return "".join(b.text for b in resp.content if b.type == "text").strip()


def set_font(run, name=BODY_FONT, size=10.5):
    run.font.name = name
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_body(doc: Document, text: str):
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        bullet = re.match(r"^[-*]\s+(.*)", line)
        p = doc.add_paragraph(style="List Bullet" if bullet else None)
        set_font(p.add_run(bullet.group(1) if bullet else line))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-i", "--input", required=True, help="조사자 메모 파일")
    ap.add_argument("-t", "--template", default="architecture_survey")
    ap.add_argument("-o", "--output", help="출력 .docx (기본: output/<입력명>_<템플릿>.docx)")
    ap.add_argument("--model", default=DEFAULT_MODEL)
    ap.add_argument("--dry-run", action="store_true", help="API 호출 없이 구조·프롬프트만 점검")
    args = ap.parse_args()

    memo = Path(args.input).read_text(encoding="utf-8")
    tpl = load_template(args.template)
    system = (ROOT / "prompts" / "_system.md").read_text(encoding="utf-8")
    out = Path(args.output) if args.output else ROOT / "output" / f"{Path(args.input).stem}_{tpl['id']}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)

    client = None
    if not args.dry_run:
        try:
            import anthropic
        except ImportError:
            sys.exit("anthropic 미설치: pip install -r requirements.txt")
        client = anthropic.Anthropic()  # ANTHROPIC_API_KEY 사용

    doc = Document()
    doc.add_heading(tpl["title"], level=0)
    note = doc.add_paragraph()
    set_font(note.add_run("※ AI 생성 초안. 연대·치수·출전은 반드시 원자료와 대조 후 확정할 것. "
                          "[확인 필요] 표시는 조사자 보완 항목."), size=9)

    for sec in tpl["sections"]:
        print(f"- {sec['heading']}", file=sys.stderr)
        user = build_prompt(sec, memo)
        text = "(dry-run) 프롬프트 길이 %d자" % len(user) if args.dry_run else generate(client, args.model, system, user)
        doc.add_heading(sec["heading"], level=1)
        add_body(doc, text)

    doc.save(out)
    print(f"저장: {out}")


if __name__ == "__main__":
    main()
